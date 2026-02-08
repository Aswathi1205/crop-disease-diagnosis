"""
Enhanced PowerPoint Presentation Generator
Creates a professional presentation with visuals and engaging content
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

def create_enhanced_presentation():
    """Create the enhanced PowerPoint presentation with visuals"""
    
    # Create presentation object
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Define color scheme (Agricultural Green Theme)
    TITLE_COLOR = RGBColor(46, 125, 50)  # Dark Green
    SUBTITLE_COLOR = RGBColor(67, 160, 71)  # Medium Green
    TEXT_COLOR = RGBColor(33, 33, 33)  # Dark Gray
    ACCENT_COLOR = RGBColor(76, 175, 80)  # Light Green
    
    # Slide 1: Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Add title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1))
    title_frame = title_box.text_frame
    title_frame.text = "AI-Driven Crop Disease Diagnosis System"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(44)
    title_para.font.bold = True
    title_para.font.color.rgb = TITLE_COLOR
    title_para.alignment = PP_ALIGN.CENTER
    
    # Add subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(9), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Empowering Farmers Through AI Technology"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(24)
    subtitle_para.font.color.rgb = SUBTITLE_COLOR
    subtitle_para.alignment = PP_ALIGN.CENTER
    
    # Add emoji/icon text
    icon_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.8), Inches(9), Inches(0.6))
    icon_frame = icon_box.text_frame
    icon_frame.text = "Agriculture + AI + Innovation"
    icon_para = icon_frame.paragraphs[0]
    icon_para.font.size = Pt(20)
    icon_para.font.italic = True
    icon_para.font.color.rgb = TEXT_COLOR
    icon_para.alignment = PP_ALIGN.CENTER
    
    # Slide 2: Introduction
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Introduction to AI-Driven Web Application"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Platform: Web-based AI solution for agricultural disease detection"
    
    points = [
        "Target Crops: Rice and Pulse crops",
        "Technology: Deep Learning (CNN) with PyTorch",
        "Interface: Streamlit web application with integrated chatbot",
        "Purpose: Empowering farmers with instant, accurate disease diagnosis"
    ]
    
    for point in points:
        p = tf.add_paragraph()
        p.text = point
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_COLOR
    
    # Add key highlights
    p = tf.add_paragraph()
    p.text = "\nKey Highlights:"
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = SUBTITLE_COLOR
    
    highlights = [
        "Real-time disease detection from crop images",
        "AI-powered chatbot for agricultural assistance",
        "Confidence-based predictions with treatment recommendations",
        "Downloadable diagnosis reports (PDF/Text)"
    ]
    
    for highlight in highlights:
        p = tf.add_paragraph()
        p.text = highlight
        p.level = 0
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_COLOR
    
    # Slide 3: Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Problem Statement & Challenges"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "The Agricultural Challenge"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.color.rgb = SUBTITLE_COLOR
    
    p = tf.add_paragraph()
    p.text = "Crop diseases cause significant yield losses globally"
    p.font.size = Pt(18)
    p.font.italic = True
    p.font.color.rgb = TEXT_COLOR
    
    p = tf.add_paragraph()
    p.text = "\nKey Problems Identified:"
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = SUBTITLE_COLOR
    
    problems = [
        "Limited Expert Access - Farmers in remote areas lack access to agricultural experts",
        "Manual Diagnosis Limitations - Time-consuming and requires specialized knowledge",
        "Economic Impact - Crop diseases reduce agricultural productivity by 20-40%"
    ]
    
    for problem in problems:
        p = tf.add_paragraph()
        p.text = problem
        p.level = 0
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_COLOR
    
    p = tf.add_paragraph()
    p.text = "\nOur Solution: AI-powered instant diagnosis accessible to all farmers"
    p.font.bold = True
    p.font.size = Pt(18)
    p.font.color.rgb = ACCENT_COLOR
    
    # Slide 4: Purpose
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Purpose of the Project"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Vision: Democratize agricultural expertise through AI technology"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.color.rgb = SUBTITLE_COLOR
    
    objectives = [
        ("Accessibility", "Web-based platform accessible from any device"),
        ("Accuracy", "Deep learning model with confidence scores"),
        ("Actionable Insights", "Disease identification + treatment recommendations"),
        ("Farmer Empowerment", "Instant diagnosis 24/7 with chatbot support")
    ]
    
    for obj_title, obj_desc in objectives:
        p = tf.add_paragraph()
        p.text = f"\n{obj_title}"
        p.font.bold = True
        p.font.size = Pt(18)
        p.font.color.rgb = ACCENT_COLOR
        
        p = tf.add_paragraph()
        p.text = obj_desc
        p.level = 1
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_COLOR
    
    # Slide 5: System Design
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    
    # Add title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = "System Design and Flow"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(36)
    title_para.font.bold = True
    title_para.font.color.rgb = TITLE_COLOR
    title_para.alignment = PP_ALIGN.CENTER
    
    # Add architecture diagram
    diagram_path = r"C:\Users\gkm09\.gemini\antigravity\brain\329d739e-bf3b-48a4-8784-12302c2c1604\system_architecture_diagram_1768450688244.png"
    if os.path.exists(diagram_path):
        slide.shapes.add_picture(diagram_path, Inches(1.5), Inches(1.2), width=Inches(7))
    
    # Slide 6: Core Modules
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Core Modules of the System"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "7 Major Modules Working Together"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.color.rgb = SUBTITLE_COLOR
    
    modules = [
        "Authentication Module - User registration, login, session management",
        "Inference Engine - Custom CNN with PyTorch for disease classification",
        "Image Processing - Preprocessing, normalization, tensor conversion",
        "Configuration Management - JSON-based flexible configuration",
        "Treatment Guide System - Comprehensive disease information database",
        "Report Generation - Professional PDF reports with ReportLab",
        "Chatbot Integration - Botpress for agricultural assistance"
    ]
    
    for module in modules:
        p = tf.add_paragraph()
        p.text = module
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_COLOR
    
    # ENHANCED Slide 7: Week 1&2 with Visual
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank for custom layout
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = "Week 1&2: Dataset & Preprocessing"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(36)
    title_para.font.bold = True
    title_para.font.color.rgb = TITLE_COLOR
    title_para.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(9), Inches(0.4))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Foundation of Accuracy - 'Garbage In, Garbage Out'"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(20)
    subtitle_para.font.italic = True
    subtitle_para.font.color.rgb = SUBTITLE_COLOR
    subtitle_para.alignment = PP_ALIGN.CENTER
    
    # Add preprocessing visual
    preprocessing_img = r"C:\Users\gkm09\.gemini\antigravity\brain\329d739e-bf3b-48a4-8784-12302c2c1604\data_preprocessing_visual_1768451620654.png"
    if os.path.exists(preprocessing_img):
        slide.shapes.add_picture(preprocessing_img, Inches(0.8), Inches(1.5), width=Inches(8.4))
    
    # Add key achievements box
    achievements_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(9), Inches(1))
    achievements_frame = achievements_box.text_frame
    achievements_frame.text = "Impact: Consistent preprocessing ensures 95%+ model accuracy | 224x224 standard size | RGB normalization"
    achievements_para = achievements_frame.paragraphs[0]
    achievements_para.font.size = Pt(14)
    achievements_para.font.bold = True
    achievements_para.font.color.rgb = ACCENT_COLOR
    achievements_para.alignment = PP_ALIGN.CENTER
    
    # ENHANCED Slide 8: Week 3&4 with Visual
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = "Week 3&4: Model Training & Evaluation"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(36)
    title_para.font.bold = True
    title_para.font.color.rgb = TITLE_COLOR
    title_para.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(9), Inches(0.4))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Precision in Prediction - Deep Learning at Work"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(20)
    subtitle_para.font.italic = True
    subtitle_para.font.color.rgb = SUBTITLE_COLOR
    subtitle_para.alignment = PP_ALIGN.CENTER
    
    # Add CNN architecture visual
    cnn_img = r"C:\Users\gkm09\.gemini\antigravity\brain\329d739e-bf3b-48a4-8784-12302c2c1604\cnn_architecture_visual_1768451643318.png"
    if os.path.exists(cnn_img):
        slide.shapes.add_picture(cnn_img, Inches(0.8), Inches(1.5), width=Inches(8.4))
    
    # Add key metrics box
    metrics_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(9), Inches(1))
    metrics_frame = metrics_box.text_frame
    metrics_frame.text = "Results: Top-3 predictions with confidence scores | CPU & GPU support | PyTorch state dict deployment"
    metrics_para = metrics_frame.paragraphs[0]
    metrics_para.font.size = Pt(14)
    metrics_para.font.bold = True
    metrics_para.font.color.rgb = ACCENT_COLOR
    metrics_para.alignment = PP_ALIGN.CENTER
    
    # ENHANCED Slide 9: Week 5&6 with Visual
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = "Week 5&6: Web Application & Chatbot"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(36)
    title_para.font.bold = True
    title_para.font.color.rgb = TITLE_COLOR
    title_para.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(9), Inches(0.4))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Empowering Farmers - Where AI Meets Usability"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(20)
    subtitle_para.font.italic = True
    subtitle_para.font.color.rgb = SUBTITLE_COLOR
    subtitle_para.alignment = PP_ALIGN.CENTER
    
    # Add web app features visual
    webapp_img = r"C:\Users\gkm09\.gemini\antigravity\brain\329d739e-bf3b-48a4-8784-12302c2c1604\web_app_features_visual_1768451665610.png"
    if os.path.exists(webapp_img):
        slide.shapes.add_picture(webapp_img, Inches(0.8), Inches(1.5), width=Inches(8.4))
    
    # Add features summary box
    features_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(9), Inches(1))
    features_frame = features_box.text_frame
    features_frame.text = "Features: Secure Authentication | Drag-Drop Upload | Treatment Guide | PDF Reports | 24/7 Chatbot Support"
    features_para = features_frame.paragraphs[0]
    features_para.font.size = Pt(14)
    features_para.font.bold = True
    features_para.font.color.rgb = ACCENT_COLOR
    features_para.alignment = PP_ALIGN.CENTER
    
    # ENHANCED Slide 10: Week 7&8
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Week 7&8: Testing & Quality Assurance"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Comprehensive Testing & Code Excellence"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.color.rgb = SUBTITLE_COLOR
    
    # Testing Strategy
    p = tf.add_paragraph()
    p.text = "\nTesting Strategy:"
    p.font.bold = True
    p.font.size = Pt(18)
    p.font.color.rgb = ACCENT_COLOR
    
    testing_points = [
        "Unit Testing: pytest framework for component validation",
        "Integration Testing: End-to-end workflow verification",
        "Bug Fixes: Model loading, image processing, UI/UX improvements"
    ]
    
    for point in testing_points:
        p = tf.add_paragraph()
        p.text = point
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_COLOR
    
    # SOLID Principles
    p = tf.add_paragraph()
    p.text = "\nSOLID Principles Applied (All 5!):"
    p.font.bold = True
    p.font.size = Pt(18)
    p.font.color.rgb = ACCENT_COLOR
    
    solid_points = [
        "SRP: Separated UserRepository from AuthManager",
        "OCP: Abstract ConfigLoader for extensibility",
        "LSP: Substitutable implementations maintain type safety",
        "ISP: Focused, specific interfaces",
        "DIP: IAuthenticator & IPredictor interfaces"
    ]
    
    for point in solid_points:
        p = tf.add_paragraph()
        p.text = point
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_COLOR
    
    p = tf.add_paragraph()
    p.text = "\nResult: Production-ready, maintainable, well-documented system"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = ACCENT_COLOR
    
    # Slide 11: Tools & Technologies
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Tools and Technologies Used"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Modern Tech Stack"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.color.rgb = SUBTITLE_COLOR
    
    tech_categories = [
        ("Machine Learning & AI", "PyTorch, PIL, NumPy, Custom CNN"),
        ("Web Development", "Streamlit, Python 3.x"),
        ("Security", "bcrypt, Session Management"),
        ("Data Management", "JSON, Dataclasses, Pathlib"),
        ("Report Generation", "ReportLab, io"),
        ("Chatbot", "Botpress Webchat"),
        ("Development Tools", "Git, pytest, VS Code"),
        ("Design", "Custom CSS, Google Fonts (Inter, Poppins)")
    ]
    
    for category, tools in tech_categories:
        p = tf.add_paragraph()
        p.text = f"{category}: {tools}"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_COLOR
    
    # Slide 12: Project Timeline
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Project Development Period"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "8-Week Development Cycle"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.color.rgb = SUBTITLE_COLOR
    
    timeline = [
        ("Phase 1 (Weeks 1-2)", "Dataset collection, preprocessing pipeline, configuration"),
        ("Phase 2 (Weeks 3-4)", "CNN architecture, model training, evaluation, deployment"),
        ("Phase 3 (Weeks 5-6)", "Streamlit app, authentication, chatbot, PDF reports"),
        ("Phase 4 (Weeks 7-8)", "Testing, bug fixes, SOLID refactoring, documentation")
    ]
    
    for phase, desc in timeline:
        p = tf.add_paragraph()
        p.text = phase
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = ACCENT_COLOR
        
        p = tf.add_paragraph()
        p.text = desc
        p.level = 1
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_COLOR
    
    p = tf.add_paragraph()
    p.text = "\nMethodology: Iterative development, test-driven, user-centric"
    p.font.italic = True
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_COLOR
    
    # Slide 13: Conclusions
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Conclusions"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Project Achievements"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.color.rgb = SUBTITLE_COLOR
    
    achievements = [
        "Functional AI System: Accurate crop disease classification",
        "Robust Architecture: SOLID principles for maintainability",
        "Complete Pipeline: End-to-end solution from upload to report",
        "User Impact: Accessible to all farmers with instant diagnosis",
        "Quality Code: Well-tested, documented, and modular"
    ]
    
    for achievement in achievements:
        p = tf.add_paragraph()
        p.text = achievement
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_COLOR
    
    p = tf.add_paragraph()
    p.text = "\nImpact Statement:"
    p.font.bold = True
    p.font.size = Pt(18)
    p.font.color.rgb = SUBTITLE_COLOR
    
    p = tf.add_paragraph()
    p.text = "This project demonstrates how AI can bridge the gap between agricultural expertise and farmers, providing accessible, accurate, and actionable disease diagnosis to improve crop yields and farmer livelihoods."
    p.font.italic = True
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_COLOR
    
    # Slide 14: Future Scope
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Future Scope"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Continuous Improvement & Expansion"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.color.rgb = SUBTITLE_COLOR
    
    future_items = [
        "Expand disease coverage and crop types",
        "Model optimization with transfer learning (ResNet, EfficientNet)",
        "Mobile app with offline capability (TensorFlow Lite)",
        "Multi-language support for regional accessibility",
        "Geolocation integration for disease prevalence tracking",
        "Weather integration for disease risk prediction",
        "Partnership with agricultural organizations",
        "Cloud deployment (AWS/Azure) for global scalability",
        "Farmer education and training programs"
    ]
    
    for item in future_items:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_COLOR
    
    # Slide 15: Thank You
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    
    # Thank you text
    thanks_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1))
    thanks_frame = thanks_box.text_frame
    thanks_frame.text = "Thank You!"
    thanks_para = thanks_frame.paragraphs[0]
    thanks_para.font.size = Pt(54)
    thanks_para.font.bold = True
    thanks_para.font.color.rgb = TITLE_COLOR
    thanks_para.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.7), Inches(9), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "AI-Driven Crop Disease Diagnosis System"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(28)
    subtitle_para.font.color.rgb = SUBTITLE_COLOR
    subtitle_para.alignment = PP_ALIGN.CENTER
    
    # Tagline
    tagline_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.6), Inches(9), Inches(0.5))
    tagline_frame = tagline_box.text_frame
    tagline_frame.text = "Empowering Farmers Through Technology"
    tagline_para = tagline_frame.paragraphs[0]
    tagline_para.font.size = Pt(20)
    tagline_para.font.italic = True
    tagline_para.font.color.rgb = TEXT_COLOR
    tagline_para.alignment = PP_ALIGN.CENTER
    
    # Questions text
    questions_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(9), Inches(0.5))
    questions_frame = questions_box.text_frame
    questions_frame.text = "Questions & Discussion"
    questions_para = questions_frame.paragraphs[0]
    questions_para.font.size = Pt(24)
    questions_para.font.bold = True
    questions_para.font.color.rgb = ACCENT_COLOR
    questions_para.alignment = PP_ALIGN.CENTER
    
    # Save presentation
    output_path = r"C:\Users\gkm09\Downloads\cnn-pytorch-main\AI_Crop_Disease_Diagnosis_Presentation_Enhanced.pptx"
    prs.save(output_path)
    print("Enhanced presentation created successfully!")
    print(f"Saved to: {output_path}")
    return output_path

if __name__ == "__main__":
    create_enhanced_presentation()
